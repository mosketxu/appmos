"""
Modelo 111 (retenciones de trabajo y profesionales): junta lo de NÓMINAS y lo de PROFESIONALES y deja el fichero para importar en
el formulario de la AEAT (Sede → modelo 111 → importar) y un Excel resumen. Pedido de Alex, 15-oct-2026. Ver ../PLAN.md.

Nóminas (no las hacemos nosotros; manda lo primero que haya):
  · txt     el fichero 111 que manda la gestoría de nóminas (diseño AEAT; trae casillas 01-06…).
  · pdf     el borrador / justificante del 111 en PDF (se leen las casillas).
  · mayor   mayor de SAGE («Fichas Mayor»): retenciones = haber de la 4751* del periodo (líneas de nóminas), percepciones = 640* del
            periodo y nº de perceptores a mano (si no, se estima por las subcuentas 465*). Vale el mayor solo de la 4751 y el completo.
Profesionales (los calculamos nosotros): listado de retenciones de SAGE (ESTLisRetencionTMP_*.xlsx), agrupado por NIF → casillas 07-09.
Si hay de las dos partes, se suman («a veces es la suma de los dos»).

  python3 prepara111.py spec.json      → escribe en spec["salida"]: 111_<NIF>_<AAAA><PP>.txt, Retenciones <PP> <AAAA><NIF>C<NOMBRE>.xlsx y resumen.json

spec = {"nif", "nombre", "ejercicio": 2026, "periodo": "09" | "3T", "salida": dir,
        "ficheros": {"txt": [...], "pdf": [...], "mayor": [...], "profesionales": [...]},
        "ajustes": {"tipo": "I|U|G|N", "iban", "perceptores_nominas", "percepciones_nominas", "retenciones_nominas",
                    "excluir_4751": [ids], "incluir_4751": [ids], "resultado_anterior", "complementaria", "justificante_anterior"}}
Diseño de registro: dr111e16v18 (AEAT; página 01 de 1000 posiciones, comprobado con el 111.txt de la gestoría de Fashion IQ).
"""
from __future__ import annotations

import datetime as dt
import json
import os
import re
import sys
import unicodedata
from collections import OrderedDict

# ---- diseño de registro (página 01) ---------------------------------------------------------------------------------------
# casillas en orden: (casilla, longitud); las de 8 son nº de perceptores, las de 17 importes con 2 decimales
CASILLAS = [('01', 8), ('02', 17), ('03', 17), ('04', 8), ('05', 17), ('06', 17), ('07', 8), ('08', 17), ('09', 17),
            ('10', 8), ('11', 17), ('12', 17), ('13', 8), ('14', 17), ('15', 17), ('16', 8), ('17', 17), ('18', 17),
            ('19', 8), ('20', 17), ('21', 17), ('22', 8), ('23', 17), ('24', 17), ('25', 8), ('26', 17), ('27', 17),
            ('28', 17), ('29', 17), ('30', 17)]
RETENCIONES = ['03', '06', '09', '12', '15', '18', '21', '24', '27']   # suman la 28
PERCEPTORES = [c for c, n in CASILLAS if n == 8]
NOMBRES = {
    '01': 'Trabajo dinerario · nº perceptores', '02': 'Trabajo dinerario · percepciones', '03': 'Trabajo dinerario · retenciones',
    '04': 'Trabajo en especie · nº perceptores', '05': 'Trabajo en especie · valor', '06': 'Trabajo en especie · ingresos a cuenta',
    '07': 'Actividades económicas dinerarias · nº perceptores', '08': 'Actividades económicas · percepciones', '09': 'Actividades económicas · retenciones',
    '10': 'Actividades económicas en especie · nº perceptores', '11': 'Actividades económicas en especie · valor', '12': 'Actividades económicas en especie · ingresos a cuenta',
    '13': 'Premios en metálico · nº perceptores', '14': 'Premios en metálico · percepciones', '15': 'Premios en metálico · retenciones',
    '16': 'Premios en especie · nº perceptores', '17': 'Premios en especie · valor', '18': 'Premios en especie · ingresos a cuenta',
    '19': 'Aprovechamientos forestales dinerarios · nº', '20': 'Aprovechamientos forestales · percepciones', '21': 'Aprovechamientos forestales · retenciones',
    '22': 'Aprovechamientos forestales en especie · nº', '23': 'Aprovechamientos forestales en especie · valor', '24': 'Aprovechamientos forestales en especie · ingresos a cuenta',
    '25': 'Derechos de imagen · nº perceptores', '26': 'Derechos de imagen · contraprestaciones', '27': 'Derechos de imagen · ingresos a cuenta',
    '28': 'Total retenciones e ingresos a cuenta', '29': 'A deducir (complementaria)', '30': 'Resultado a ingresar',
}


class ErrorDatos(Exception):
    pass


def limpia(s: str) -> str:
    """Solo letras, números y blancos (los An del diseño): «FASHION IQ BCN, S.L.» → «FASHION IQ BCN SL»."""
    s = unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode()
    s = re.sub(r'[^A-Za-z0-9 ]+', '', s.replace(',', ' ').replace('.', ''))
    return re.sub(r'\s+', ' ', s).strip().upper()


def periodo_meses(periodo: str) -> list[int]:
    if re.fullmatch(r'[1-4]T', periodo):
        t = int(periodo[0])
        return [3 * t - 2, 3 * t - 1, 3 * t]
    if re.fullmatch(r'(0[1-9]|1[0-2])', periodo):
        return [int(periodo)]
    raise ErrorDatos(f'periodo no válido: {periodo} (01…12 o 1T…4T)')


def num(v) -> float:
    if v is None or v == '':
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace('€', '').replace(' ', '')
    neg = s.startswith('-') or s.startswith('(')
    s = s.strip('-()')
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    try:
        x = float(s)
    except ValueError:
        return 0.0
    return -x if neg else x


# ---- lectura de ficheros ---------------------------------------------------------------------------------------------------
def filas_xlsx(ruta: str) -> list[list]:
    import openpyxl
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        for fila in ws.iter_rows(values_only=True):
            out.append(list(fila))
    return out


def cabecera(filas: list[list], obligatorias: list[str]) -> tuple[int, dict]:
    for i, f in enumerate(filas[:60]):
        nombres = {str(v).strip().lower(): j for j, v in enumerate(f) if v is not None}
        if all(o.lower() in nombres for o in obligatorias):
            return i, nombres
    raise ErrorDatos('no encuentro la cabecera (' + ', '.join(obligatorias) + ')')


def formato_fechas(valores: list) -> str:
    """'dmy' o 'mdy' para las fechas en texto de una columna (SAGE las exporta de las dos maneras)."""
    for v in valores:
        if isinstance(v, str):
            m = re.match(r'(\d{1,2})/(\d{1,2})/(\d{4})', v.strip())
            if m and int(m.group(1)) > 12:
                return 'dmy'
            if m and int(m.group(2)) > 12:
                return 'mdy'
    return 'dmy'


def fecha(v, fmt: str) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):   # número de serie de Excel
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
    m = re.match(r'(\d{1,2})/(\d{1,2})/(\d{4})', str(v or '').strip())
    if not m:
        return None
    a, b, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    d, mm = (a, b) if fmt == 'dmy' else (b, a)
    try:
        return dt.date(y, mm, d)
    except ValueError:
        return None


def lee_txt(ruta: str) -> dict:
    """Fichero 111 de la AEAT (el que manda la gestoría): casillas, NIF, periodo, tipo, IBAN."""
    s = open(ruta, encoding='latin-1').read()
    i = s.find('<T11101000>')
    if i < 0:
        raise ErrorDatos(f'«{os.path.basename(ruta)}» no es un fichero del modelo 111 (no tiene la página <T11101000>)')
    p = s[i:i + 1000]
    if len(p) < 1000:
        raise ErrorDatos(f'«{os.path.basename(ruta)}»: la página 01 está cortada')
    out = {'tipo': p[12], 'nif': p[13:22].strip(), 'nombre': p[22:82].strip(), 'ejercicio': p[102:106], 'periodo': p[106:108],
           'iban': p[552:586].strip(), 'complementaria': p[537] == 'X', 'justificante_anterior': p[538:551].strip(), 'casillas': {}}
    pos = 108
    for c, n in CASILLAS:
        v = p[pos:pos + n]
        pos += n
        neg = v.startswith('N')
        v = v.lstrip('N')
        if not v.strip().isdigit():
            raise ErrorDatos(f'«{os.path.basename(ruta)}»: la casilla {c} no es un número («{v}»)')
        x = int(v) if n == 8 else int(v) / 100
        out['casillas'][c] = -x if neg else x
    return out


IMPORTE = re.compile(r'^\(?-?\d{1,3}(?:\.\d{3})*,\d\d\)?$|^\(?-?\d+,\d\d\)?$')
ENTERO = re.compile(r'^\d{1,3}(?:\.\d{3})*$|^\d+$')


def lee_pdf(ruta: str) -> dict:
    """Borrador o justificante del 111 de la AEAT: casillas por posición (tres columnas por fila: perceptores | percepciones | retenciones)."""
    import pymupdf
    doc = pymupdf.open(ruta)
    cas = {}
    out = {'casillas': cas}
    for pag in doc:
        ws = pag.get_text('words')
        for w in ws:   # «Ejercicio ... 2026   Período ..... 2T» (por posición: el orden del texto del PDF no es el de la página)
            if w[4] == 'Ejercicio' and 'ejercicio' not in out:
                fila = sorted((x for x in ws if abs(x[1] - w[1]) < 4 and x[0] > w[0]), key=lambda x: x[0])
                a = next((x[4] for x in fila if re.fullmatch(r'\d{4}', x[4])), None)
                pp = next((x[4] for x in fila if re.fullmatch(r'(0[1-9]|1[0-2]|[1-4]T)', x[4])), None)
                if a:
                    out['ejercicio'], out['periodo'] = a, pp
        m = re.search(r'\b([A-Z]\d{7}[A-Z0-9]|\d{8}[A-Z])\b', ' '.join(w[4] for w in ws if 190 < w[1] < 230))
        if m and 'nif' not in out:
            out['nif'] = m.group(1)
        etiquetas = [w for w in ws if re.fullmatch(r'\d{2}', w[4]) and w[4] in NOMBRES and w[0] > 255]
        # fuera la fila de la fórmula de la 28 «( 03 + 06 + … + 27 )»: una fila de etiquetas de verdad tiene como mucho tres
        etiquetas = [e for e in etiquetas if e[4] in ('28', '29', '30') or sum(1 for o in etiquetas if abs(o[1] - e[1]) < 3) <= 3]
        for e in etiquetas:
            c = e[4]
            if c in cas:
                continue
            x0, y0 = e[0], e[1]
            if c in ('28', '29', '30'):
                cerca = [w for w in ws if abs(w[1] - y0) < 6 and w[0] > x0 + 10 and IMPORTE.match(w[4])]
                if cerca:
                    cas[c] = num(max(cerca, key=lambda w: w[0])[4])
                continue
            # columna de la casilla: perceptores (x≈264), percepciones (x≈346), retenciones (x≈461); el valor va en la misma fila o justo encima
            col = 0 if x0 < 300 else (1 if x0 < 420 else 2)
            lim = [(255, 346), (346, 461), (461, 600)][col]
            cand = [w for w in ws if y0 - 14 <= w[1] <= y0 + 4 and lim[0] < w[0] < lim[1] and w is not e
                    and (ENTERO.match(w[4]) if col == 0 else IMPORTE.match(w[4]))]
            cand = [w for w in cand if not (re.fullmatch(r'\d{2}', w[4]) and w[4] in NOMBRES and abs(w[1] - y0) < 2)]
            if cand:
                w = max(cand, key=lambda w: w[1])
                cas[c] = int(w[4].replace('.', '')) if col == 0 else num(w[4])
    if not cas:
        raise ErrorDatos(f'«{os.path.basename(ruta)}»: no encuentro casillas del modelo 111 en el PDF')
    for c, n in CASILLAS:
        cas.setdefault(c, 0 if n == 8 else 0.0)
    return out


def lee_mayor(ruta: str, meses: list[int], ejercicio: int) -> dict:
    """Mayor de SAGE: líneas de la 4751* y de la 640*/641* y subcuentas 465* del periodo."""
    filas = filas_xlsx(ruta)
    i, h = cabecera(filas, ['Fecha asiento', 'Código cuenta', 'Debe', 'Haber'])
    ic, ifc, idb, ihb = h['código cuenta'], h['fecha asiento'], h['debe'], h['haber']
    icom, inum, icp = h.get('comentario'), h.get('número de asiento'), h.get('contrapartida')
    icif = h.get('cif/dni')
    fmt = formato_fechas([f[ifc] for f in filas[i + 1:] if len(f) > ifc])
    out = {'4751': [], '640': [], '465': set()}
    base = os.path.basename(ruta)
    for k, f in enumerate(filas[i + 1:], start=i + 2):
        if len(f) <= max(ic, ifc, idb, ihb) or f[ic] is None:
            continue
        cta = str(f[ic]).strip().split('.')[0]
        d = fecha(f[ifc], fmt)
        if not d or d.year != ejercicio or d.month not in meses:
            continue
        com = str(f[icom] or '').replace('\n', ' ').strip() if icom is not None else ''
        debe, haber = num(f[idb]), num(f[ihb])
        lin = {'id': f'{base}#{k}', 'fecha': d.strftime('%d/%m/%Y'), 'asiento': str(f[inum] or '') if inum is not None else '',
               'cuenta': cta, 'contrapartida': str(f[icp] or '') if icp is not None else '', 'comentario': com, 'debe': round(debe, 2), 'haber': round(haber, 2)}
        if cta.startswith('4751'):
            out['4751'].append(lin)
        elif cta.startswith('640') or cta.startswith('641'):
            out['640'].append(lin)
        elif cta.startswith('465') and haber > 0:
            out['465'].add((cta, str(f[icif] or '') if icif is not None else ''))
    return out


def lee_profesionales(ruta: str, meses: list[int], ejercicio: int, avisos: list) -> list[dict]:
    """Listado de retenciones de SAGE (ESTLisRetencionTMP_*): una fila por factura; solo las recibidas (R) que no son alquileres."""
    filas = filas_xlsx(ruta)
    i, h = cabecera(filas, ['CifDni', 'BaseRetencion', 'ImporteRetencion'])
    g = lambda f, k: f[h[k.lower()]] if k.lower() in h and h[k.lower()] < len(f) else None
    fmt = formato_fechas([g(f, 'FechaFactura') for f in filas[i + 1:]])
    out = []
    base = os.path.basename(ruta)
    for f in filas[i + 1:]:
        nif = str(g(f, 'CifDni') or '').strip()
        if not nif and not g(f, 'ImporteRetencion'):
            continue
        tipo = str(g(f, 'TipoFactura') or 'R').strip().upper()
        nombre = str(g(f, 'ClienteProveedor') or '').strip()
        clave = str(g(f, 'DescripcionClave') or '').strip()
        r = {'nombre': nombre, 'nif': nif, 'fecha': '', 'factura': str(g(f, 'SuFacturaNo') or g(f, 'Factura') or ''),
             'base': round(num(g(f, 'BaseRetencion')), 2), 'pct': num(g(f, '%Retencion')), 'retencion': round(num(g(f, 'ImporteRetencion')), 2),
             'clave': clave, 'fichero': base}
        d = fecha(g(f, 'FechaFactura'), fmt)
        r['fecha'] = d.strftime('%d/%m/%Y') if d else str(g(f, 'FechaFactura') or '')
        if tipo != 'R':
            avisos.append(f'Profesionales: fuera «{nombre}» ({r["factura"]}): es una factura emitida (retención que nos practican, no va en el 111).')
            continue
        if re.search(r'arrend|alquil', clave, re.I):
            avisos.append(f'Profesionales: fuera «{nombre}» ({r["factura"]}): clave «{clave}» (va en el 115).')
            continue
        if d and (d.year != ejercicio or d.month not in meses):
            avisos.append(f'Profesionales: «{nombre}» ({r["factura"]}) tiene fecha {r["fecha"]}, fuera del periodo; se incluye (el listado de SAGE manda).')
        out.append(r)
    return out


def nif_clave(v) -> str:
    """NIF para comparar: mayúsculas, sin espacios ni guiones y sin el «ES» delante."""
    n = re.sub(r'[^A-Z0-9]', '', str(v or '').upper())
    return n[2:] if re.fullmatch(r'ES[A-Z0-9]{9}', n) else n


def lee_cp_proveedores(ruta: str) -> dict:
    """Listado de proveedores de SAGE (lisProveedores*.xlsx): {NIF: código postal} (columnas CifDni y CodigoPostal)."""
    filas = filas_xlsx(ruta)
    i, h = cabecera(filas, ['CifDni', 'CodigoPostal'])
    out = {}
    for f in filas[i + 1:]:
        nif = nif_clave(f[h['cifdni']] if h['cifdni'] < len(f) else '')
        cp = f[h['codigopostal']] if h['codigopostal'] < len(f) else ''
        cp = str(int(cp)) if isinstance(cp, float) and cp == int(cp) else str(cp or '').strip()
        if nif and cp:
            out.setdefault(nif, cp.zfill(5) if cp.isdigit() and len(cp) < 5 else cp)
    return out


# ---- cálculo ----------------------------------------------------------------------------------------------------------------
def prepara(spec: dict) -> dict:
    ej = int(spec['ejercicio'])
    per = str(spec['periodo'])
    meses = periodo_meses(per)
    aj = spec.get('ajustes') or {}
    fic = spec.get('ficheros') or {}
    avisos, errores = [], []
    nif = str(spec.get('nif') or '').strip().upper()
    nombre = limpia(spec.get('nombre') or '')[:60]
    if not re.fullmatch(r'[A-Z0-9]{9}', nif):
        errores.append(f'NIF de la empresa no válido: «{nif}».')

    cas = OrderedDict((c, 0 if n == 8 else 0.0) for c, n in CASILLAS)
    fuentes = []
    nominas = {'origen': '', 'casillas': {}, 'lineas_4751': [], 'lineas_640': []}

    # --- nóminas: txt > pdf > mayor
    txts = [lee_txt(r) for r in fic.get('txt', [])]
    pdfs = []
    for r in fic.get('pdf', []):
        try:
            pdfs.append(lee_pdf(r))
        except ErrorDatos as e:
            errores.append(str(e))
    for t, r in zip(txts, fic.get('txt', [])):
        if t['nif'] and t['nif'] != nif:
            avisos.append(f'El 111.txt «{os.path.basename(r)}» es del NIF {t["nif"]}, no de {nif}.')
        if t['ejercicio'] != str(ej) or t['periodo'] != per:
            avisos.append(f'El 111.txt «{os.path.basename(r)}» es de {t["periodo"]}/{t["ejercicio"]}, no de {per}/{ej}.')
    if txts:
        for t in txts:
            for c, _ in CASILLAS[:27]:
                nominas['casillas'][c] = nominas['casillas'].get(c, 0) + t['casillas'][c]
        nominas['origen'] = 'txt de la gestoría (' + ', '.join(os.path.basename(r) for r in fic['txt']) + ')'
        if not aj.get('iban') and txts[0]['iban']:
            aj.setdefault('iban_txt', txts[0]['iban'])
        if pdfs:
            for c in ('01', '02', '03', '04', '05', '06'):
                if abs(pdfs[0]['casillas'].get(c, 0) - nominas['casillas'].get(c, 0)) > 0.005:
                    avisos.append(f'El PDF y el txt de nóminas no coinciden en la casilla {c}: PDF {pdfs[0]["casillas"].get(c)} · txt {nominas["casillas"].get(c)}. Manda el txt.')
    elif pdfs:
        for p in pdfs:
            for c, _ in CASILLAS[:27]:
                nominas['casillas'][c] = nominas['casillas'].get(c, 0) + p['casillas'].get(c, 0)
            if p.get('nif') and p['nif'] != nif:
                avisos.append(f'El PDF de nóminas es del NIF {p["nif"]}, no de {nif}.')
            if p.get('periodo') and (p['periodo'] != per or p.get('ejercicio') != str(ej)):
                avisos.append(f'El PDF de nóminas es de {p["periodo"]}/{p.get("ejercicio")}, no de {per}/{ej}.')
        nominas['origen'] = 'PDF de nóminas (' + ', '.join(os.path.basename(r) for r in fic['pdf']) + ')'

    lineas_4751, lineas_640, subc_465 = [], [], set()
    for r in fic.get('mayor', []):
        try:
            m = lee_mayor(r, meses, ej)
        except ErrorDatos as e:
            errores.append(f'Mayor «{os.path.basename(r)}»: {e}')
            continue
        lineas_4751 += m['4751']
        lineas_640 += m['640']
        subc_465 |= m['465']
    # una línea repetida en dos mayores (el de la 4751 y el completo) cuenta una vez
    vistos, unicas = set(), []
    for l in lineas_4751:
        k = (l['fecha'], l['asiento'], l['cuenta'], l['debe'], l['haber'], l['comentario'])
        if k not in vistos:
            vistos.add(k)
            unicas.append(l)
    lineas_4751 = unicas
    vistos, unicas = set(), []
    for l in lineas_640:
        k = (l['fecha'], l['asiento'], l['cuenta'], l['debe'], l['haber'])
        if k not in vistos:
            vistos.add(k)
            unicas.append(l)
    lineas_640 = unicas
    excl, incl = set(aj.get('excluir_4751') or []), set(aj.get('incluir_4751') or [])
    for l in lineas_4751:
        es_nomina = bool(re.search(r'n[oó]min', l['comentario'], re.I))
        es_pago = l['debe'] > 0 or re.match(r'apertura', l['comentario'], re.I)
        defecto = es_nomina and not es_pago and l['haber'] > 0
        l['motivo'] = '' if defecto else ('pago / apertura' if es_pago else ('no parece de nóminas' if l['haber'] > 0 else 'cargo'))
        l['incluida'] = (defecto or l['id'] in incl) and l['id'] not in excl
    nominas['lineas_4751'] = lineas_4751
    nominas['lineas_640'] = lineas_640

    if not nominas['origen'] and lineas_4751:
        ret = round(sum(l['haber'] for l in lineas_4751 if l['incluida']), 2)
        perc = round(sum(l['debe'] - l['haber'] for l in lineas_640 if l['cuenta'].startswith('640')), 2)
        if any(l['cuenta'].startswith('641') for l in lineas_640):
            avisos.append('Hay movimientos en la 641 (indemnizaciones) en el periodo: no se suman a las percepciones; revísalo.')
        nominas['casillas'] = {'01': len(subc_465) or 0, '02': perc, '03': ret}
        nominas['origen'] = 'mayor de SAGE (4751 → retenciones' + (', 640 → percepciones' if lineas_640 else '') + ')'
        if not lineas_640:
            avisos.append('El mayor no trae la 640 (sueldos) del periodo: las percepciones de nóminas (casilla 02) hay que escribirlas o subir el mayor completo.')
        if subc_465 and aj.get('perceptores_nominas') in (None, ''):
            avisos.append(f'Nº de perceptores de nóminas estimado por las subcuentas 465 del periodo: {len(subc_465)}. Compruébalo.')
        fuera = [l for l in lineas_4751 if not l['incluida'] and l['haber'] > 0 and l['motivo'] != 'pago / apertura']
        if fuera:
            avisos.append(f'{len(fuera)} línea(s) de la 4751 no parecen de nóminas y no se han sumado; márcalas si lo son.')
    elif nominas['origen'] and lineas_4751:
        ret = round(sum(l['haber'] for l in lineas_4751 if l['incluida']), 2)
        if abs(ret - nominas['casillas'].get('03', 0)) > 0.005:
            avisos.append(f'Las retenciones de nóminas del mayor (4751: {ret:,.2f}) no coinciden con las del {nominas["origen"].split(" (")[0]} ({nominas["casillas"].get("03", 0):,.2f}).')

    # ajustes a mano de nóminas
    for clave, c in (('perceptores_nominas', '01'), ('percepciones_nominas', '02'), ('retenciones_nominas', '03')):
        v = aj.get(clave)
        if v not in (None, ''):
            nominas['casillas'][c] = int(num(v)) if c == '01' else round(num(v), 2)
            nominas.setdefault('a_mano', []).append(c)
    for c, v in nominas['casillas'].items():
        cas[c] += v
    if nominas['origen']:
        fuentes.append('Nóminas: ' + nominas['origen'])
        if cas['03'] and not cas['01']:
            errores.append('Falta el nº de perceptores de nóminas (casilla 01).')
        if cas['03'] and not cas['02']:
            errores.append('Faltan las percepciones de nóminas (casilla 02).')

    # --- profesionales
    prof = []
    for r in fic.get('profesionales', []):
        try:
            prof += lee_profesionales(r, meses, ej, avisos)
        except ErrorDatos as e:
            errores.append(f'Profesionales «{os.path.basename(r)}»: {e}')
    # código postal de cada profesional: del listado de proveedores de SAGE (fichero base del cliente)
    cps = {}
    for r in fic.get('proveedores', []):
        try:
            cps.update(lee_cp_proveedores(r))
        except ErrorDatos as e:
            avisos.append(f'Proveedores «{os.path.basename(r)}»: {e}; sin código postal.')
    for p in prof:
        p['cp'] = cps.get(nif_clave(p['nif']), '')
    if prof and not fic.get('proveedores'):
        avisos.append('Sin listado de proveedores de SAGE: el Excel sale sin código postal de los profesionales (súbelo como fichero base «Proveedores»).')
    elif prof:
        faltan = sorted({p['nombre'] or p['nif'] for p in prof if not p['cp']})
        if faltan:
            avisos.append('Sin código postal en el listado de proveedores: ' + ', '.join(faltan[:10]) + (f' y {len(faltan) - 10} más' if len(faltan) > 10 else '') + '.')
    grupos = OrderedDict()
    for p in prof:
        g = grupos.setdefault(p['nif'] or p['nombre'], {'nif': p['nif'], 'nombre': p['nombre'], 'cp': p['cp'], 'base': 0.0, 'retencion': 0.0, 'facturas': 0})
        g['base'] = round(g['base'] + p['base'], 2)
        g['retencion'] = round(g['retencion'] + p['retencion'], 2)
        g['facturas'] += 1
    if prof:
        if nominas['casillas'].get('09'):
            avisos.append('La parte de nóminas ya trae actividades económicas (casillas 07-09) y se le suman los profesionales: comprueba que no estén repetidos.')
        cas['07'] += len(grupos)
        cas['08'] = round(cas['08'] + sum(g['base'] for g in grupos.values()), 2)
        cas['09'] = round(cas['09'] + sum(g['retencion'] for g in grupos.values()), 2)
        fuentes.append('Profesionales: listado de SAGE (' + ', '.join(os.path.basename(r) for r in fic['profesionales']) + f') · {len(prof)} factura(s), {len(grupos)} perceptor(es)')

    # --- liquidación
    cas['28'] = round(sum(cas[c] for c in RETENCIONES), 2)
    cas['29'] = round(num(aj.get('resultado_anterior')), 2)
    cas['30'] = round(cas['28'] - cas['29'], 2)
    if not fuentes:
        errores.append('Sube al menos un fichero (nóminas o profesionales).')
    iban = re.sub(r'\s', '', str(aj.get('iban') or aj.get('iban_txt') or '')).upper()
    tipo = str(aj.get('tipo') or '').upper() or ('U' if iban else 'I')
    if cas['30'] <= 0:
        tipo = 'N'
    if tipo == 'U' and not re.fullmatch(r'[A-Z]{2}\d{2}[A-Z0-9]{10,30}', iban):
        errores.append('Domiciliación (U) sin IBAN válido.')
    if tipo not in ('I', 'U', 'G', 'N'):
        errores.append(f'Tipo de declaración no válido: {tipo}.')
    compl = bool(aj.get('complementaria'))
    just = re.sub(r'\D', '', str(aj.get('justificante_anterior') or ''))
    if compl and len(just) != 13:
        errores.append('Complementaria: falta el nº de justificante de la declaración anterior (13 cifras).')

    return {'nif': nif, 'nombre': nombre, 'ejercicio': ej, 'periodo': per, 'tipo': tipo, 'iban': iban if tipo == 'U' else '',
            'complementaria': compl, 'justificante_anterior': just if compl else '',
            'casillas': dict(cas), 'nominas': nominas, 'profesionales': prof, 'perceptores_prof': list(grupos.values()),
            'fuentes': fuentes, 'avisos': avisos, 'errores': errores}


# ---- fichero AEAT ------------------------------------------------------------------------------------------------------------
def _num(v, n: int, signo: bool = False) -> str:
    if n == 8:
        s = str(int(v)).rjust(8, '0')
    else:
        cent = int(round(v * 100))
        if cent < 0:
            if not signo:
                raise ErrorDatos(f'importe negativo en un campo sin signo: {v}')
            s = 'N' + str(-cent).rjust(n - 1, '0')
        else:
            s = str(cent).rjust(n, '0')
    if len(s) != n:
        raise ErrorDatos(f'importe demasiado grande: {v}')
    return s


def genera(d: dict) -> str:
    per, ej = d['periodo'], str(d['ejercicio'])
    cab = f'<T1110{ej}{per}0000>' + '<AUX>' + ' ' * 300 + '</AUX>'
    p = '<T11101000>' + ' ' + d['tipo'] + d['nif'].ljust(9)[:9] + d['nombre'].ljust(60)[:60] + ' ' * 20 + ej + per
    for c, n in CASILLAS:
        p += _num(d['casillas'][c], n, signo=(c == '29'))
    p += ('X' if d['complementaria'] else ' ') + d['justificante_anterior'].ljust(13)[:13] + ' ' + d['iban'].ljust(34)[:34]
    p += ' ' * 389 + ' ' * 13 + '</T11101000>'
    if len(p) != 1000:
        raise ErrorDatos(f'la página 01 mide {len(p)} y debería medir 1000')
    if len(cab) != 328:
        raise ErrorDatos(f'la cabecera mide {len(cab)}')
    return cab + p + f'</T1110{ej}{per}0000>'


# ---- Excel resumen ------------------------------------------------------------------------------------------------------------
def excel(d: dict, ruta: str) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Retenciones'
    neg = Font(bold=True)
    gris = PatternFill('solid', start_color='FFE5E7EB')
    fino = Side(style='thin', color='FF9CA3AF')
    caja = Border(top=fino, bottom=fino, left=fino, right=fino)
    eur = '#,##0.00'
    r = 1
    ws.cell(r, 1, f'Modelo 111 · {d["nombre"]} · {d["nif"]} · {d["periodo"]} {d["ejercicio"]}').font = Font(bold=True, size=12)
    r += 2

    def cab(fila, textos):
        for j, t in enumerate(textos, 1):
            c = ws.cell(fila, j, t)
            c.font, c.fill, c.border = neg, gris, caja

    # Profesionales (como el resumen de SAGE que se usaba: perceptor, NIF, fecha, factura, base, %, retención)
    if d['profesionales']:
        ws.cell(r, 1, 'PROFESIONALES (actividades económicas · casillas 07-09)').font = neg
        r += 1
        cab(r, ['ClienteProveedor', 'CifDni', 'FechaFactura', 'SuFacturaNo', 'BaseRetencion', '%Retencion', 'ImporteRetencion', 'CodigoPostal'])
        r += 1
        for p in d['profesionales']:
            for j, v in enumerate([p['nombre'], p['nif'], p['fecha'], p['factura'], p['base'], p['pct'], p['retencion'], p.get('cp', '')], 1):
                c = ws.cell(r, j, v)
                c.border = caja
                if j in (5, 7):
                    c.number_format = eur
            r += 1
        ws.cell(r, 4, 'Total').font = neg
        for j, k in ((5, 'base'), (7, 'retencion')):
            c = ws.cell(r, j, round(sum(p[k] for p in d['profesionales']), 2))
            c.font, c.number_format = neg, eur
        r += 2
        if len(d['perceptores_prof']) != len(d['profesionales']):
            ws.cell(r, 1, 'Por perceptor (NIF)').font = neg
            r += 1
            cab(r, ['Perceptor', 'NIF', 'Facturas', '', 'Base', '', 'Retención', 'CodigoPostal'])
            r += 1
            for g in d['perceptores_prof']:
                for j, v in ((1, g['nombre']), (2, g['nif']), (3, g['facturas']), (5, g['base']), (7, g['retencion']), (8, g.get('cp', ''))):
                    c = ws.cell(r, j, v)
                    if j in (5, 7):
                        c.number_format = eur
                r += 1
            r += 1

    # Nóminas
    nom = d['nominas']
    if nom['origen']:
        ws.cell(r, 1, 'NÓMINAS (rendimientos del trabajo · casillas 01-06) · ' + nom['origen']).font = neg
        r += 1
        lin = [l for l in nom['lineas_4751'] if l['incluida']]
        if lin and nom['origen'].startswith('mayor'):
            cab(r, ['Fecha asiento', 'Número de asiento', 'Código cuenta', 'Contrapartida', 'Comentario', 'Debe', 'Haber'])
            r += 1
            for l in lin:
                for j, v in enumerate([l['fecha'], l['asiento'], l['cuenta'], l['contrapartida'], l['comentario'], l['debe'], l['haber']], 1):
                    c = ws.cell(r, j, v)
                    c.border = caja
                    if j in (6, 7):
                        c.number_format = eur
                r += 1
            c = ws.cell(r, 7, round(sum(l['haber'] for l in lin), 2))
            c.font, c.number_format = neg, eur
            r += 2
        else:
            for c_ in ('01', '02', '03', '04', '05', '06'):
                if nom['casillas'].get(c_):
                    ws.cell(r, 1, f'[{c_}] {NOMBRES[c_]}')
                    c = ws.cell(r, 7, nom['casillas'][c_])
                    c.number_format = '0' if c_ in PERCEPTORES else eur
                    r += 1
            r += 1

    # Casillas del 111
    ws.cell(r, 1, 'MODELO 111 · casillas').font = neg
    r += 1
    cab(r, ['Casilla', 'Concepto', '', '', '', '', 'Importe'])
    r += 1
    for c_, _ in CASILLAS:
        v = d['casillas'][c_]
        if not v and c_ not in ('28', '30'):
            continue
        ws.cell(r, 1, c_)
        ws.cell(r, 2, NOMBRES[c_])
        c = ws.cell(r, 7, v)
        c.number_format = '0' if c_ in PERCEPTORES else eur
        if c_ in ('28', '30'):
            for j in (1, 2, 7):
                ws.cell(r, j).font = neg
        r += 1
    r += 1
    ws.cell(r, 1, 'Tipo de declaración')
    ws.cell(r, 2, {'I': 'Ingreso', 'U': 'Domiciliación', 'G': 'Ingreso en CCT', 'N': 'Negativa'}.get(d['tipo'], d['tipo']) + (f' · {d["iban"]}' if d['iban'] else ''))
    for a in d['avisos']:
        r += 1
        ws.cell(r, 1, '⚠ ' + a).font = Font(color='FFB45309')
    for k, w in zip('ABCDEFG', (38, 14, 13, 16, 14, 11, 16)):
        ws.column_dimensions[k].width = w
    for fila in ws.iter_rows():
        for c in fila:
            c.alignment = Alignment(vertical='top')
    wb.save(ruta)


def nombre_excel(d: dict) -> str:
    return f'Retenciones {d["periodo"]} {d["ejercicio"]}{d["nif"]}C{d["nombre"]}.xlsx'


def main() -> None:
    spec = json.load(open(sys.argv[1], encoding='utf-8'))
    salida = spec['salida']
    os.makedirs(salida, exist_ok=True)
    for f in os.listdir(salida):
        os.remove(os.path.join(salida, f))
    try:
        d = prepara(spec)
    except ErrorDatos as e:
        d = {'errores': [str(e)], 'avisos': []}
    if not d['errores']:
        try:
            txt = genera(d)
            d['fichero'] = f'111_{d["nif"]}_{d["ejercicio"]}{d["periodo"]}.txt'
            with open(os.path.join(salida, d['fichero']), 'w', encoding='latin-1', newline='') as fh:
                fh.write(txt)
        except ErrorDatos as e:
            d['errores'].append(str(e))
    if 'casillas' in d:
        try:
            d['excel'] = nombre_excel(d)
            excel(d, os.path.join(salida, d['excel']))
        except Exception as e:   # el Excel es un extra: no para el fichero
            d['avisos'].append(f'No se ha podido hacer el Excel resumen: {e}')
            d.pop('excel', None)
    json.dump(d, open(os.path.join(salida, 'resumen.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1, default=str)
    print(json.dumps({'ok': not d['errores'], 'errores': d['errores']}, ensure_ascii=False))


if __name__ == '__main__':
    main()
